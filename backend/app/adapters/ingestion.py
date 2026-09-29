"""Bounded in-memory parsers; uploads are never retained."""
import csv
import io
import json
import zipfile
from pathlib import Path
from xml.etree.ElementTree import ParseError

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.config import Settings, settings


def validate_header(header, limits):
    if not header or len(header) > limits.max_columns:
        raise ValueError(f'Expected 1 to {limits.max_columns} columns.')
    if any(not isinstance(name, str) or not name.strip() for name in header):
        raise ValueError('Column names must be nonempty strings.')
    if len(set(header)) != len(header):
        raise ValueError('Column names must be unique.')


def parse_upload(filename: str, content: bytes, limits: Settings = settings) -> pd.DataFrame:
    if not content or len(content) > limits.max_upload_bytes:
        raise ValueError(f'Upload must contain 1 to {limits.max_upload_bytes} bytes.')
    extension = Path(filename).suffix.lower()
    try:
        if extension == '.csv':
            text = content.decode('utf-8-sig')
            if '\x00' in text:
                raise ValueError('CSV must be UTF-8 text.')
            reader = csv.reader(io.StringIO(text), strict=True)
            header = next(reader)
            validate_header(header, limits)
            count = 0
            for row in reader:
                if not row:
                    continue
                if len(row) != len(header):
                    raise ValueError('CSV row width differs from header.')
                count += 1
                if count * len(header) > limits.max_cells:
                    raise ValueError('Dataset cell limit exceeded.')
                if count > limits.max_rows:
                    raise ValueError('Row limit exceeded.')
            frame = pd.read_csv(io.StringIO(text), nrows=limits.max_rows + 1)
        elif extension == '.xlsx':
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > limits.max_xlsx_expanded_bytes:
                    raise ValueError('Expanded XLSX size limit exceeded.')
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                sheet = workbook.worksheets[0]
                if sheet.max_column and sheet.max_column > limits.max_columns:
                    raise ValueError('Column limit exceeded.')
                if sheet.max_row and sheet.max_row > limits.max_rows + 1:
                    raise ValueError('Row limit exceeded (including blank worksheet rows).')
                iterator = sheet.iter_rows(values_only=True)
                header = list(next(iterator))
                validate_header(header, limits)
                rows = []
                for row in iterator:
                    if all(value is None for value in row):
                        continue
                    rows.append(row)
                    if len(rows) * len(header) > limits.max_cells:
                        raise ValueError('Dataset cell limit exceeded.')
                    if len(rows) > limits.max_rows:
                        raise ValueError('Row limit exceeded.')
                frame = pd.DataFrame(rows, columns=header)
            finally:
                workbook.close()
        elif extension == '.json':
            def reject_constant(value):
                raise ValueError('Non-finite JSON numbers are unsupported.')
            data = json.loads(content.decode('utf-8-sig'), parse_constant=reject_constant)
            if isinstance(data, list):
                if not data or not all(isinstance(row, dict) for row in data):
                    raise ValueError('Expected a nonempty array of records.')
                if len(data) > limits.max_rows:
                    raise ValueError('Row limit exceeded.')
                names = set()
                for row in data:
                    names.update(row)
                    if len(names) > limits.max_columns:
                        raise ValueError('Column limit exceeded.')
                if len(names) * len(data) > limits.max_cells:
                    raise ValueError('Dataset cell limit exceeded.')
            elif isinstance(data, dict):
                if not data or not all(isinstance(values, list) for values in data.values()):
                    raise ValueError('Tabular JSON must map columns to arrays.')
                if max(map(len, data.values())) > limits.max_rows:
                    raise ValueError('Row limit exceeded.')
                validate_header(list(data), limits)
                if len(data) * max(map(len, data.values())) > limits.max_cells:
                    raise ValueError('Dataset cell limit exceeded.')
            else:
                raise ValueError('Expected JSON records or column arrays.')
            frame = pd.DataFrame(data)
            if frame.map(lambda value: isinstance(value, (dict, list))).any().any():
                raise ValueError('Nested JSON cells are unsupported.')
        else:
            raise ValueError('Supported extensions: .csv, .xlsx, .json.')
        validate_header(list(frame.columns), limits)
        if frame.size > limits.max_cells:
            raise ValueError('Dataset cell limit exceeded.')
        if frame.empty or len(frame) > limits.max_rows:
            raise ValueError(f'Expected 1 to {limits.max_rows} rows.')
        return frame
    except (ValueError, UnicodeError, csv.Error, zipfile.BadZipFile, StopIteration,
            KeyError, TypeError, OSError, ParseError, InvalidFileException, IndexError) as exc:
        raise ValueError(f'Invalid {extension or "file"} upload: {exc}') from None
