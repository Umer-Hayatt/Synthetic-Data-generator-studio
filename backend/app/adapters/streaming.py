"""Disk-backed batch readers and bounded reservoir profiles; no full-source frame."""
import json
from pathlib import Path
import zipfile
import ijson
import numpy as np
import pandas as pd
import pyarrow.csv as pc
import pyarrow.parquet as pq
from openpyxl import load_workbook
from app.adapters.ingestion import parse_upload, validate_header
from app.core.config import settings
from app.core.profiling import fit_spec


def batches(path, format, batch_rows=None):
    size = batch_rows or settings.batch_rows
    if size < 1:
        raise ValueError('Batch size must be positive.')
    if format == 'csv':
        with pc.open_csv(path, read_options=pc.ReadOptions(block_size=1024**2, use_threads=False)) as reader:
            for batch in reader:
                for offset in range(0, batch.num_rows, size):
                    yield batch.slice(offset, size).to_pandas()
    elif format == 'parquet':
        with pq.ParquetFile(path) as reader:
            if len(reader.schema_arrow) > settings.max_columns:
                raise ValueError('Column limit exceeded.')
            # Avoid decompression amplification beyond this local adapter's budget.
            if any(reader.metadata.row_group(i).total_byte_size > 64 * 1024**2 for i in range(reader.num_row_groups)):
                raise ValueError('Re-export Parquet with row groups below 64 MiB.')
            for batch in reader.iter_batches(batch_size=size, use_threads=False):
                yield batch.to_pandas()
    elif format in ('json', 'jsonl'):
        with Path(path).open('rb') as stream:
            if format == 'json':
                first = stream.read(4096).lstrip()[:1]
                stream.seek(0)
                if first != b'[':
                    if Path(path).stat().st_size > settings.max_upload_bytes:
                        raise ValueError('Large column-array JSON should be converted to JSONL or Parquet.')
                    yield parse_upload('data.json', stream.read())
                    return
                iterator = ijson.items(stream, 'item', use_float=True)
            else:
                def lines():
                    while line := stream.readline(1024**2 + 1):
                        if len(line) > 1024**2:
                            raise ValueError('JSONL record exceeds 1 MiB; use smaller records.')
                        if line.strip():
                            yield json.loads(line)
                iterator = lines()
            rows = []
            for row in iterator:
                if not isinstance(row, dict) or any(isinstance(v, (dict, list)) for v in row.values()):
                    raise ValueError('Expected flat JSON records.')
                validate_header(list(row), settings)
                rows.append(row)
                if len(rows) == size:
                    yield pd.DataFrame(rows)
                    rows = []
            if rows:
                yield pd.DataFrame(rows)
    elif format == 'xlsx':
        with zipfile.ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > settings.max_xlsx_expanded_bytes:
                raise ValueError('Excel expansion exceeds deployment limit; convert to CSV or Parquet.')
        source = Path(path).open('rb')
        book = load_workbook(source, read_only=True, data_only=True)
        try:
            sheet = book.worksheets[0]
            if sheet.max_column and sheet.max_column > settings.max_columns:
                raise ValueError('Excel column limit exceeded; export a smaller sheet.')
            iterator = sheet.iter_rows(values_only=True)
            header = list(next(iterator))
            validate_header(header, settings)
            rows = []
            for row in iterator:
                if all(value is None for value in row):
                    continue
                rows.append(row)
                if len(rows) == size:
                    yield pd.DataFrame(rows, columns=header)
                    rows = []
            if rows:
                yield pd.DataFrame(rows, columns=header)
        finally:
            book.close()
            source.close()
    else:
        raise ValueError('Supported formats: csv, json, jsonl, xlsx, parquet.')


def profile_source(path, format, progress=lambda stage, fraction: None, sample_rows=None, include_sample=False):
    budget = min(sample_rows or settings.profile_rows, settings.max_rows)
    rng = np.random.default_rng(42)
    reservoir, priorities = pd.DataFrame(), np.array([])
    count, nulls, columns = 0, {}, None
    for frame in batches(path, format):
        progress('profiling', .2)
        validate_header(list(frame.columns), settings)
        if columns is None:
            columns = list(frame.columns)
            nulls = dict.fromkeys(columns, 0)
        if list(frame.columns) != columns:
            raise ValueError('Column layout changes between batches; normalize records first.')
        count += len(frame)
        if count > settings.job_max_rows:
            raise ValueError('Deployment profiling row limit exceeded.')
        if frame.memory_usage(deep=True).sum() > 64 * 1024**2:
            raise ValueError('Batch exceeds local memory budget; reduce BATCH_ROWS or cell size.')
        for name in columns:
            nulls[name] += int(frame[name].isna().sum())
        reservoir = pd.concat([reservoir, frame], ignore_index=True)
        priorities = np.concatenate([priorities, rng.random(len(frame))])
        if len(reservoir) > budget:
            indices = np.argpartition(priorities, budget-1)[:budget]
            reservoir, priorities = reservoir.iloc[indices].reset_index(drop=True), priorities[indices]
    if not count:
        raise ValueError('Source contains no records.')
    spec = fit_spec(reservoir)
    profile = {'row_count': count, 'sample_rows': len(reservoir), 'approximate': count > len(reservoir),
            'null_counts': nulls, 'spec': spec.model_dump(mode='json'),
            'generation_row_count_note': 'Review row_count; the fitted specification uses the bounded sample size.'}
    return (profile, reservoir) if include_sample else profile
